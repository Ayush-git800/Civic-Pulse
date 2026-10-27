import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen slides (13.33 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    BG_DARK = RGBColor(15, 23, 42)        # #0F172A (Deep Slate)
    CARD_BG = RGBColor(30, 41, 59)        # #1E293B (Card Slate)
    CARD_BORDER = RGBColor(51, 65, 85)    # #334155
    TEXT_LIGHT = RGBColor(248, 250, 252)  # #F8FAFC
    TEXT_MUTED = RGBColor(148, 163, 184)  # #94A3B8
    INDIGO = RGBColor(99, 102, 241)       # #6366F1
    VIOLET = RGBColor(139, 92, 246)       # #8B5CF6
    EMERALD = RGBColor(16, 185, 129)      # #10B981
    AMBER = RGBColor(245, 158, 11)        # #F59E0B
    ROSE = RGBColor(244, 63, 94)          # #F43F5E

    def set_slide_background(slide):
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = BG_DARK
        bg_shape.line.color.rgb = BG_DARK
        return bg_shape

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
        return card

    # ==========================================
    # SLIDE 1: Title & Team
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Hackathon Badge Pill
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.8), Inches(5.2), Inches(0.5))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(49, 46, 129)
    badge.line.color.rgb = INDIGO
    tf = badge.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "⚡ Hacktoberfect Hackathon 2026"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RGBColor(199, 210, 254)
    p.alignment = PP_ALIGN.CENTER

    # Project Title
    tx_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.4), Inches(11.3), Inches(1.8))
    tf = tx_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "CivicPulse"
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT

    p2 = tf.add_paragraph()
    p2.text = "AI-Powered Civic Hazard Triage & Hyperlocal Home Utility Marketplace"
    p2.font.size = Pt(22)
    p2.font.bold = False
    p2.font.color.rgb = INDIGO

    p3 = tf.add_paragraph()
    p3.text = "Built on Google's Gemma 4 Multimodal Architecture | Target Tracks: Open Source & Best Use of Gemma"
    p3.font.size = Pt(14)
    p3.font.color.rgb = TEXT_MUTED

    # Team Header Label
    th_box = s1.shapes.add_textbox(Inches(1.0), Inches(3.6), Inches(11.3), Inches(0.4))
    tf = th_box.text_frame
    p = tf.paragraphs[0]
    p.text = "PROUDLY DEVELOPED BY TEAM: VERSION CONTROLLERS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    # Team Member Cards (3 columns, no individual roles)
    members = [
        ("Ayush Chauhan", INDIGO),
        ("Rishav Ranjan", VIOLET),
        ("Hardik Kumar", EMERALD),
    ]

    card_w = 3.5
    gap = 0.4
    for i, (name, accent) in enumerate(members):
        c_left = 1.0 + i * (card_w + gap)
        add_card(s1, c_left, 4.2, card_w, 2.2)

        # Accent top bar
        bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(c_left), Inches(4.2), Inches(card_w), Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.color.rgb = accent

        tb = s1.shapes.add_textbox(Inches(c_left + 0.2), Inches(4.7), Inches(card_w - 0.4), Inches(1.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = name
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT

        p_t = tf.add_paragraph()
        p_t.text = "Team Version Controllers"
        p_t.font.size = Pt(13)
        p_t.font.color.rgb = accent

    # ==========================================
    # SLIDE 2: Problem & Dual Solution
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)

    # Slide Title
    st = s2.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.3), Inches(1.1))
    tf = st.text_frame
    p = tf.paragraphs[0]
    p.text = "The Dual Crisis & The CivicPulse Solution"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT
    p2 = tf.add_paragraph()
    p2.text = "Solving public infrastructure neglect and private home repair exploitation through multimodal AI."
    p2.font.size = Pt(15)
    p2.font.color.rgb = TEXT_MUTED

    # Column 1: Problem Space
    add_card(s2, 1.0, 1.8, 5.4, 5.0)
    tb1 = s2.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(5.0), Inches(4.6))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "⚠️ THE PROBLEM"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ROSE

    points_p = [
        ("Bureaucratic Black Hole:", "Citizens encounter dangerous potholes, open junction boxes, or toxic waste, but don't know whether to notify PWD, Municipal Corp, or Electricity Board."),
        ("Jurisdiction Mismatch:", "Complaints are dismissed when citizens report National Highway or Railway hazards to city councils, causing weeks of inaction."),
        ("Home Utility Price Gouging:", "Homeowners lack technical knowledge, resulting in arbitrary inflated quotes from technicians and opaque material costs."),
        ("High Aggregator Commissions:", "Commercial platforms extract up to 25% from daily-wage technicians, driving up costs for consumers.")
    ]
    for title, desc in points_p:
        p_item = tf1.add_paragraph()
        p_item.text = f"• {title} "
        p_item.font.bold = True
        p_item.font.size = Pt(12)
        p_item.font.color.rgb = TEXT_LIGHT
        run = p_item.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # Column 2: CivicPulse Solution
    add_card(s2, 6.9, 1.8, 5.4, 5.0)
    tb2 = s2.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.0), Inches(4.6))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "✨ THE CIVICPULSE SOLUTION"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    points_s = [
        ("Gemma 4 Forensic Multimodal Triage:", "Citizens upload a single photo; Gemma 4 automatically classifies severity, municipal department, and remediation steps in seconds."),
        ("Intelligent Jurisdictional Routing:", "Distinguishes local city mandates from central authorities (auto-generates direct escalation links to CPGRAMS, NHAI, RailMadad)."),
        ("AI Fair Price Estimation:", "Diagnoses household failures (plumbing, electrical, carpentry) and attaches an objective price bracket for labor and parts."),
        ("Transparent Bargaining Engine:", "Technicians bid directly, homeowners counter-offer, and both lock fair rates with zero middleman commissions.")
    ]
    for title, desc in points_s:
        p_item = tf2.add_paragraph()
        p_item.text = f"✔ {title} "
        p_item.font.bold = True
        p_item.font.size = Pt(12)
        p_item.font.color.rgb = TEXT_LIGHT
        run = p_item.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 3: Key Features & Gemma 4 Power
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)

    st = s3.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.3), Inches(1.1))
    tf = st.text_frame
    p = tf.paragraphs[0]
    p.text = "How Gemma 4 Powers CivicPulse"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT
    p2 = tf.add_paragraph()
    p2.text = "Harnessing Google's open multimodal model for civic auditing and transparent local commerce."
    p2.font.size = Pt(15)
    p2.font.color.rgb = TEXT_MUTED

    features = [
        ("🔬 1. Zero-Shot Forensic Audit", "Gemma 4 processes raw image bytes to assess hazard depth, physical risk, and structural failure with municipal engineer accuracy.", INDIGO),
        ("🏛️ 2. Autonomous Escalation", "Detects central vs city jurisdictions. Non-local issues auto-generate direct links to CPGRAMS, NHAI, and RailMadad portals.", VIOLET),
        ("💰 3. Fair Price Bracketing", "Analyzes broken appliances or plumbing fixtures to output standard Indian market cost brackets, stopping consumer price gouging.", EMERALD),
        ("🤝 4. Worker-Customer Bidding", "Technicians place bids, homeowners counter-bid, and both lock in deals on an open, community-driven marketplace.", AMBER),
        ("🤖 5. PulseBot Public Copilot", "Real-time AI assistant answering civic governance procedures, emergency helplines, and project inquiries with live streaming.", INDIGO),
        ("⚡ 6. Resilient Fallback Cascade", "99.9% uptime architecture cascading from Gemma 4 to Gemini Flash models if quota thresholds are exceeded.", ROSE),
    ]

    for idx, (ft_title, ft_desc, accent) in enumerate(features):
        row = idx // 3
        col = idx % 3
        c_left = 1.0 + col * (3.6 + 0.25)
        c_top = 1.8 + row * (2.4 + 0.25)

        add_card(s3, c_left, c_top, 3.6, 2.4)
        bar = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(c_left), Inches(c_top), Inches(3.6), Inches(0.08))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.color.rgb = accent

        tb = s3.shapes.add_textbox(Inches(c_left + 0.15), Inches(c_top + 0.2), Inches(3.3), Inches(2.0))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = ft_title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT

        p_d = tf.add_paragraph()
        p_d.text = ft_desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 4: Architecture & Engineering Rigor
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)

    st = s4.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.3), Inches(1.1))
    tf = st.text_frame
    p = tf.paragraphs[0]
    p.text = "System Architecture & Engineering Rigor"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT
    p2 = tf.add_paragraph()
    p2.text = "Production-grade, zero-bloat stack engineered for resilience, speed, and cross-platform accessibility."
    p2.font.size = Pt(15)
    p2.font.color.rgb = TEXT_MUTED

    stack_cards = [
        ("🎨 FRONTEND LAYER", "Vanilla JS (ES6+) + Tailwind CSS", [
            "Zero JavaScript framework overhead (No bloated node_modules)",
            "Instant DOM rendering & interactive masonry grid",
            "Animated laser scanner visual inspection effect",
            "Edge-to-edge forensic evidence lightbox modal"
        ], INDIGO),
        ("⚙️ BACKEND & APIS", "FastAPI (Python 3.9+) + Uvicorn", [
            "Asynchronous high-concurrency request processing",
            "Pydantic v2 strict payload validation schemas",
            "Multi-part binary image stream ingestion",
            "Full CORS compliance & static asset serving"
        ], VIOLET),
        ("🧠 AI ORCHESTRATION", "Gemma 4 (26B-A4B-IT) + GenAI SDK", [
            "Native multimodal image-byte inspection",
            "Strict JSON schema generation with sanitizers",
            "Resilient fallback cascade (Gemini 3.5/3.8 Flash)",
            "Zero-latency failure isolation & dynamic recovery"
        ], EMERALD),
        ("🗄️ PERSISTENCE", "SQLite 3 (WAL Mode & Relational Integrity)", [
            "Thread-safe connection pool with mutex locks",
            "Write-Ahead Logging (WAL) for concurrent read/write",
            "Relational foreign keys for posts, jobs & bids",
            "Auto-seeding test fixtures for instant demoing"
        ], AMBER)
    ]

    for idx, (layer_title, tech, bullets, accent) in enumerate(stack_cards):
        c_left = 1.0 + idx * (2.7 + 0.17)
        add_card(s4, c_left, 1.8, 2.7, 5.0)

        bar = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(c_left), Inches(1.8), Inches(2.7), Inches(0.1))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.color.rgb = accent

        tb = s4.shapes.add_textbox(Inches(c_left + 0.15), Inches(2.0), Inches(2.4), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = layer_title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = accent

        p_t = tf.add_paragraph()
        p_t.text = tech
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_LIGHT

        for b in bullets:
            p_b = tf.add_paragraph()
            p_b.text = f"• {b}"
            p_b.font.size = Pt(10)
            p_b.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 5: Future Roadmap & Vision
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)

    st = s5.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.3), Inches(1.1))
    tf = st.text_frame
    p = tf.paragraphs[0]
    p.text = "Future Roadmap: Scaling CivicPulse"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT
    p2 = tf.add_paragraph()
    p2.text = "High-impact features planned for our production deployment and urban expansion."
    p2.font.size = Pt(15)
    p2.font.color.rgb = TEXT_MUTED

    roadmap = [
        ("🚀 PHASE 1: DIRECT MUNICIPAL APIS", "Direct Integration with ULBs & Smart Cities Mission", "Connect directly into Urban Local Bodies (ULBs) and Smart Cities Mission grievance databases for auto-ticketing without human clerk delays.", INDIGO),
        ("📱 PHASE 2: ON-DEVICE EDGE GEMMA", "Quantized Offline Gemma on Mobile", "Deploy quantized Gemma 2B directly on low-spec smartphones using MediaPipe, allowing citizens in rural or low-connectivity areas to audit issues offline.", VIOLET),
        ("🗣️ PHASE 3: VOICE-FIRST & VERNACULAR BOT", "Multilingual Indian Voice Redressal", "Integrate Bhashini open-source speech models to allow uneducated citizens and daily-wage workers to report hazards via voice in 12+ Indian languages.", EMERALD),
        ("💳 PHASE 4: UPI ESCROW & CONTRACTS", "Milestone-Based Smart Payments", "Introduce escrow payments: homeowner locks payment in UPI escrow; funds release only when Gemma 4 verifies the repair completion photo.", AMBER),
        ("🔮 PHASE 5: PREDICTIVE INFRASTRUCTURE", "Pre-Accident Deterioration Forecasting", "Use time-series satellite and street camera feeds to predict which roads will collapse during monsoon before accidents happen.", ROSE),
    ]

    r_w = 11.3
    r_h = 0.88
    r_gap = 0.16
    for idx, (phase, title, detail, accent) in enumerate(roadmap):
        r_top = 1.8 + idx * (r_h + r_gap)
        add_card(s5, 1.0, r_top, r_w, r_h)

        # Left color pill
        lp = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(r_top), Inches(0.12), Inches(r_h))
        lp.fill.solid()
        lp.fill.fore_color.rgb = accent
        lp.line.color.rgb = accent

        tb = s5.shapes.add_textbox(Inches(1.25), Inches(r_top + 0.08), Inches(10.8), Inches(r_h - 0.16))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"{phase}  —  "
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = accent

        r_title = p.add_run()
        r_title.text = title
        r_title.font.bold = True
        r_title.font.size = Pt(13)
        r_title.font.color.rgb = TEXT_LIGHT

        p_det = tf.add_paragraph()
        p_det.text = detail
        p_det.font.size = Pt(11)
        p_det.font.color.rgb = TEXT_MUTED

    # Output file
    output_filename = "CivicPulse_Hackathon_Pitch_Deck.pptx"
    prs.save(output_filename)
    print(f"Presentation saved successfully as '{output_filename}'.")

if __name__ == "__main__":
    create_deck()
